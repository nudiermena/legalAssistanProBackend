from fastapi import APIRouter, HTTPException, Body, Depends, Request, BackgroundTasks
from fastapi.responses import FileResponse
from typing import List, Optional, Dict, Any, Union, Literal
from pydantic import BaseModel, Field, ConfigDict, validator
from datetime import datetime, timezone
from agents.document_drafting_agent import draft_legal_document, draft_custom_document
from config.colombian_compliance import ColombianLegalFramework
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO
import markdown
import os as os_module
import shutil
import tempfile
from utils.r2_storage import upload_document_to_r2
from endpoints.auth import get_current_user
from utils.security_validators import validate_document_request, sanitize_text_input, validate_legal_document_request
from utils.security_monitor import SecurityMonitor
import logging
from bs4 import BeautifulSoup

# Security logger
security_logger = logging.getLogger("security")

# Create a single router without prefix - we'll specify full paths
router = APIRouter(tags=["documents"])

class Party(BaseModel):
    name: str
    role: str = Field(default="Parte")
    details: Optional[str] = None

class CustomParty(BaseModel):
    name: str = Field(..., description="Nombre/Razón social")
    identification: str = Field(..., description="NIT/CC/CE")
    role: str = Field(..., description="Rol en el documento")

class CustomDocumentRequest(BaseModel):
    document_type: str = Field(..., description="Tipo de documento")
    description: str = Field(..., description="Descripción detallada del documento")
    parties: List[CustomParty] = Field(..., description="Partes involucradas")
    key_points: str = Field(..., description="Puntos clave a incluir")
    industry: str = Field(..., description="Industria")
    jurisdiction: str = Field(..., description="Jurisdicción aplicable")
    complexity: Literal["simple", "standard", "complex"] = Field(..., description="Complejidad del documento")
    
    model_config = ConfigDict(
        arbitrary_types_allowed=True,
        json_schema_extra={
            "example": {
                "document_type": "Contrato de colaboración",
                "description": "Contrato para colaboración entre empresas tecnológicas",
                "parties": [
                    {
                        "name": "Empresa ABC S.A.S.",
                        "identification": "900123456-7",
                        "role": "Contratante"
                    },
                    {
                        "name": "Empresa XYZ Ltda.",
                        "identification": "800987654-3",
                        "role": "Contratista"
                    }
                ],
                "key_points": "Confidencialidad\nPropiedad intelectual\nTérminos de pago\nDuración del contrato",
                "industry": "technology",
                "jurisdiction": "federal",
                "complexity": "standard"
            }
        }
    )

class DocumentDraftingRequest(BaseModel):
    document_type: str = Field(..., description="Tipo de documento")
    jurisdiction: str = Field(default="Colombia", description="Jurisdicción")
    key_requirements: List[str] = Field(default_factory=list, description="Requisitos clave")
    parties: List[Dict[str, str]] = Field(..., description="Partes")
    legal_terms: Optional[List[str]] = Field(default_factory=list, description="Términos jurídicos")
    data_processing: Optional[Dict[str, str]] = Field(default_factory=dict, description="Tratamiento de datos")

    model_config = ConfigDict(
        arbitrary_types_allowed=True,
        json_schema_extra={
            "example": {
                "document_type": "contrato_laboral",
                "key_requirements": ["Término indefinido", "Período de prueba"],
                "parties": [
                    {"name": "Empresa ABC", "role": "Empleador"},
                    {"name": "Juan Pérez", "role": "Trabajador"}
                ]
            }
        }
    )

class DocumentDraftingResponse(BaseModel):
    document: Dict[str, Union[str, List[str], bool]] = Field(..., description="Documento generado")
    document_type: str = Field(..., description="Tipo de documento")
    jurisdiction: str = Field(..., description="Jurisdicción")
    key_requirements: List[str] = Field(..., description="Requisitos")
    parties: List[Dict[str, str]] = Field(..., description="Partes")
    colombian_compliance: Dict[str, Any] = Field(..., description="Cumplimiento")
    legal_terms: Optional[Dict[str, str]] = None
    data_processing: Optional[Dict[str, str]] = Field(
        None,
        description="Configuración de tratamiento de datos",
        example={
            "category": "personal",
            "retention_period": "365 días",
            "consent_valid": "true",
            "requirements": "Requisitos estándar de protección de datos"
        }
    )
    document_url: Optional[str] = Field(None, description="URL del documento en R2")
    timestamp: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    complexity: Optional[str] = Field(None, description="Complejidad del documento")
    industry: Optional[str] = Field(None, description="Industria")

    model_config = ConfigDict(
        arbitrary_types_allowed=True,
        json_schema_extra={
            "example": {
                "document": {
                    "content": "Entre los suscritos...",
                    "sections": ["Partes", "Objeto", "Obligaciones"],
                    "markdown": True
                },
                "document_url": "https://document-generator.r2.dev/contrato_laboral_123.docx"
            }
        }
    )

class HybridDocumentRequest(BaseModel):
    """Hybrid document request that can handle both formats"""
    document_type: str = Field(..., description="Tipo de documento")
    description: Optional[str] = Field(None, description="Descripción detallada del documento")
    jurisdiction: str = Field(default="Colombia", description="Jurisdicción")
    key_requirements: Optional[List[str]] = Field(default_factory=list, description="Requisitos clave")
    key_points: Optional[str] = Field(None, description="Puntos clave a incluir")
    parties: List[Dict[str, str]] = Field(..., description="Partes")
    legal_terms: Optional[List[str]] = Field(default_factory=list, description="Términos jurídicos")
    data_processing: Optional[Dict[str, str]] = Field(default_factory=dict, description="Tratamiento de datos")
    industry: Optional[str] = Field(None, description="Industria")
    complexity: Optional[Literal["simple", "standard", "complex"]] = Field(None, description="Complejidad del documento")

    model_config = ConfigDict(
        arbitrary_types_allowed=True,
        json_schema_extra={
            "example": {
                "document_type": "contrato_laboral",
                "description": "Contrato de trabajo por prestación de servicios",
                "jurisdiction": "medellin",
                "key_requirements": ["Término indefinido", "Período de prueba"],
                "parties": [
                    {"name": "Nudier Mena", "role": "Contratante"},
                    {"name": "Karen Mena", "role": "Contratista"}
                ],
                "legal_terms": [],
                "data_processing": {"type": "personal", "consent": "explicit"},
                "industry": "other",
                "complexity": "standard"
            }
        }
    )

def create_word_document(content: str, document_type: str) -> BytesIO:
    """Create a Word document from HTML/Markdown content with advanced formatting."""
    # Convert markdown to HTML if not already HTML
    html_content = markdown.markdown(content)
    soup = BeautifulSoup(html_content, "html.parser")
    doc = Document()
    # Set up the document style
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)
    # Add title
    title = doc.add_heading(document_type.replace('_', ' ').title(), 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_html_element(element, parent=None):
        if element.name in ["p", None]:
            # Paragraph or plain text
            text = element.get_text() if element.name else str(element)
            if text.strip():
                doc.add_paragraph(text.strip())
        elif element.name in ["strong", "b"]:
            p = doc.add_paragraph()
            run = p.add_run(element.get_text())
            run.bold = True
        elif element.name in ["em", "i"]:
            p = doc.add_paragraph()
            run = p.add_run(element.get_text())
            run.italic = True
        elif element.name in ["ul", "ol"]:
            style = "List Bullet" if element.name == "ul" else "List Number"
            for li in element.find_all("li", recursive=False):
                doc.add_paragraph(li.get_text(), style=style)
        elif element.name in ["h1", "h2", "h3", "h4", "h5", "h6"]:
            level = int(element.name[1])
            doc.add_heading(element.get_text(), level=level)
        elif element.name == "hr":
            doc.add_paragraph("\u2015" * 20)  # horizontal line
        elif element.name == "table":
            rows = element.find_all("tr")
            if rows:
                cols = rows[0].find_all(["td", "th"])
                table = doc.add_table(rows=len(rows), cols=len(cols))
                for i, row in enumerate(rows):
                    cells = row.find_all(["td", "th"])
                    for j, cell in enumerate(cells):
                        table.cell(i, j).text = cell.get_text()
        elif element.name == "img":
            # TODO: Download and insert image if src is valid and accessible
            doc.add_paragraph("[Imagen: {}]".format(element.get("src", "")))
        else:
            # Fallback: add text content
            if hasattr(element, 'get_text'):
                text = element.get_text()
                if text.strip():
                    doc.add_paragraph(text.strip())

    # Add all top-level elements
    for elem in soup.body or soup:
        if hasattr(elem, 'name') or str(elem).strip():
            add_html_element(elem)

    # Save to BytesIO
    doc_io = BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io

# Removed HTML document drafting interface route

@router.post(
    "/api/v1/draft-custom-document",
    response_model=DocumentDraftingResponse,
    summary="Generación de Documentos Personalizados",
    description="Genera documentos jurídicos personalizados según especificaciones detalladas del usuario"
)
async def draft_custom_document_endpoint(
    request: CustomDocumentRequest = Body(
        ...,
        description="Parámetros para la generación del documento personalizado"
    ),
    current_user: Dict[str, Any] = Depends(get_current_user),
    http_request: Request = None
):
    """Genera documentos jurídicos personalizados con cumplimiento normativo colombiano"""
    try:
        # Security logging
        client_ip = http_request.client.host if http_request and http_request.client else "unknown"
        security_logger.info(f"Custom document request from user {current_user.get('id', 'unknown')} at {client_ip}")
        
        # Input validation and sanitization
        validated_request = validate_document_request({
            "document_type": sanitize_text_input(request.document_type),
            "description": sanitize_text_input(request.description),
            "parties": [
                {
                    "name": sanitize_text_input(party.name),
                    "identification": sanitize_text_input(party.identification),
                    "role": sanitize_text_input(party.role)
                }
                for party in request.parties
            ],
            "key_points": sanitize_text_input(request.key_points),
            "industry": request.industry,
            "jurisdiction": request.jurisdiction,
            "complexity": request.complexity
        })
        
        # Convert parties to the format expected by the agent
        party_data = [
            {
                "name": party["name"],
                "identification": party["identification"],
                "role": party["role"]
            }
            for party in validated_request["parties"]
        ]
        
        # Convert key_points string to list
        key_points_list = [point.strip() for point in validated_request["key_points"].split('\n') if point.strip()]
        
        # Security monitoring
        SecurityMonitor.log_document_generation(
            user_id=current_user.get('id', 'unknown'),
            document_type=validated_request["document_type"],
            complexity=validated_request["complexity"],
            industry=validated_request["industry"],
            client_ip=client_ip
        )
        
        result = await draft_custom_document(
            document_type=validated_request["document_type"],
            description=validated_request["description"],
            parties=party_data,
            key_points=key_points_list,
            industry=validated_request["industry"],
            jurisdiction=validated_request["jurisdiction"],
            complexity=validated_request["complexity"]
        )
        
        # Create Word document
        doc_io = create_word_document(result["document"]["content"], validated_request["document_type"])
        
        # Upload to R2
        document_url = await upload_document_to_r2(doc_io, validated_request["document_type"])
        
        response_data = {
            "document": result["document"],
            "document_type": result["document_type"],
            "jurisdiction": result["jurisdiction"],
            "key_requirements": result["key_requirements"],
            "parties": result["parties"],
            "colombian_compliance": result["colombian_compliance"],
            "legal_terms": result.get("legal_terms"),
            "data_processing": result.get("data_processing"),
            "document_url": document_url,
            "timestamp": datetime.now(timezone.utc),
            "complexity": validated_request["complexity"],
            "industry": validated_request["industry"]
        }
        
        # Success logging
        security_logger.info(f"Custom document generated successfully for user {current_user.get('id', 'unknown')}")
        
        return DocumentDraftingResponse(**response_data)
    
    except Exception as e:
        # Error logging
        security_logger.error(f"Error generating custom document for user {current_user.get('id', 'unknown')}: {str(e)}")
        SecurityMonitor.log_security_event(
            event_type="document_generation_error",
            user_id=current_user.get('id', 'unknown'),
            details={"error": str(e), "document_type": request.document_type if request else "unknown"}
        )
        
        raise HTTPException(
            status_code=500,
            detail={
                "error": str(e),
                "message": "Error en la generación del documento personalizado",
                "timestamp": datetime.now(timezone.utc).isoformat()
            }
        )

@router.post(
    "/api/v1/draft-document",
    response_model=DocumentDraftingResponse,
    summary="Generación de Documentos Jurídicos",
    description="Genera documentos jurídicos según el ordenamiento jurídico colombiano"
)
async def draft_document_endpoint(
    request: DocumentDraftingRequest = Body(
        ...,
        description="Parámetros para la generación del documento"
    ),
    current_user: Dict[str, Any] = Depends(get_current_user),
    http_request: Request = None
):
    """Genera documentos jurídicos con cumplimiento normativo colombiano"""
    try:
        # Security logging
        client_ip = http_request.client.host if http_request and http_request.client else "unknown"
        security_logger.info(f"Document request from user {current_user.get('id', 'unknown')} at {client_ip}")
        
        # Input validation and sanitization
        validated_request = validate_legal_document_request({
            "document_type": request.document_type,
            "jurisdiction": request.jurisdiction,
            "key_requirements": request.key_requirements,
            "parties": request.parties,
            "legal_terms": request.legal_terms,
            "data_processing": request.data_processing
        })
        
        party_names = [party.get("name", "") for party in validated_request["parties"]]
        
        # Security monitoring
        SecurityMonitor.log_document_generation(
            user_id=current_user.get('id', 'unknown'),
            document_type=validated_request["document_type"],
            complexity="standard",  # Default for original endpoint
            industry="general",  # Default for original endpoint
            client_ip=client_ip
        )
        
        result = await draft_legal_document(
            document_type=validated_request["document_type"],
            jurisdiction=validated_request["jurisdiction"],
            key_requirements=validated_request["key_requirements"],
            parties=party_names,
            legal_terms=validated_request["legal_terms"],
            data_processing=validated_request["data_processing"]
        )
        
        # Create Word document
        doc_io = create_word_document(result["document"]["content"], validated_request["document_type"])
        
        # Upload to R2
        document_url = await upload_document_to_r2(doc_io, validated_request["document_type"])
        
        response_data = {
            "document": result["document"],
            "document_type": result["document_type"],
            "jurisdiction": result["jurisdiction"],
            "key_requirements": result["key_requirements"],
            "parties": result["parties"],
            "colombian_compliance": result["colombian_compliance"],
            "legal_terms": result.get("legal_terms"),
            "data_processing": result.get("data_processing"),
            "document_url": document_url,
            "timestamp": datetime.now(timezone.utc)
        }
        
        # Success logging
        security_logger.info(f"Document generated successfully for user {current_user.get('id', 'unknown')}")
        
        return DocumentDraftingResponse(**response_data)
    
    except Exception as e:
        # Error logging
        security_logger.error(f"Error generating document for user {current_user.get('id', 'unknown')}: {str(e)}")
        SecurityMonitor.log_security_event(
            event_type="document_generation_error",
            user_id=current_user.get('id', 'unknown'),
            details={"error": str(e), "document_type": request.document_type if request else "unknown"}
        )
        
        raise HTTPException(
            status_code=500,
            detail={
                "error": str(e),
                "message": "Error en la generación del documento",
                "timestamp": datetime.now(timezone.utc).isoformat()
            }
        )

@router.get(
    "/document-drafting/templates",
    response_model=Dict[str, List[str]],
    summary="Plantillas Disponibles",
    description="Obtiene las plantillas de documentos jurídicos disponibles"
)
async def get_document_templates(current_user: Dict[str, Any] = Depends(get_current_user)):
    """Retorna las plantillas de documentos disponibles"""
    return {
        "contratos_laborales": [
            "Término indefinido",
            "Término fijo",
            "Obra o labor",
            "Aprendizaje"
        ],
        "contratos_comerciales": [
            "Compraventa",
            "Arrendamiento",
            "Prestación de servicios",
            "Distribución"
        ],
        "documentos_societarios": [
            "Acta de constitución",
            "Reforma estatutaria",
            "Acta de asamblea",
            "Poder especial"
        ],
        "documentos_procesales": [
            "Demanda",
            "Contestación",
            "Recurso de apelación",
            "Tutela"
        ]
    }

@router.get(
    "/document-drafting/industries",
    response_model=Dict[str, str],
    summary="Industrias Disponibles",
    description="Obtiene las industrias disponibles para documentos personalizados"
)
async def get_industries(current_user: Dict[str, Any] = Depends(get_current_user)):
    """Retorna las industrias disponibles"""
    return {
        "technology": "Tecnología",
        "healthcare": "Salud",
        "finance": "Finanzas",
        "real-estate": "Bienes raíces",
        "retail": "Comercio",
        "agriculture": "Agricultura",
        "oil-gas": "Petróleo y Gas",
        "education": "Educación",
        "construction": "Construcción",
        "other": "Otra"
    }

@router.get(
    "/document-drafting/jurisdictions",
    response_model=Dict[str, str],
    summary="Jurisdicciones Disponibles",
    description="Obtiene las jurisdicciones disponibles para documentos personalizados"
)
async def get_jurisdictions(current_user: Dict[str, Any] = Depends(get_current_user)):
    """Retorna las jurisdicciones disponibles"""
    return {
        "federal": "Nacional",
        "bogota": "Bogotá D.C.",
        "medellin": "Medellín",
        "cali": "Cali",
        "barranquilla": "Barranquilla",
        "bucaramanga": "Bucaramanga",
        "cartagena": "Cartagena",
        "pereira": "Pereira",
        "other": "Otra"
    }

@router.get(
    "/document-drafting/complexity-levels",
    response_model=Dict[str, str],
    summary="Niveles de Complejidad",
    description="Obtiene los niveles de complejidad disponibles para documentos personalizados"
)
async def get_complexity_levels(current_user: Dict[str, Any] = Depends(get_current_user)):
    """Retorna los niveles de complejidad disponibles"""
    return {
        "simple": "Simple - Términos básicos",
        "standard": "Estándar - Nivel de detalle normal",
        "complex": "Complejo - Altamente detallado"
    }

def get_legal_framework(doc_type: str, requirements: List[str]) -> Dict[str, List[str]]:
    """Obtiene marco jurídico según tipo de documento y requisitos"""
    return {
        "normativa": ["Norma 1", "Norma 2"],
        "jurisprudencia": ["Sentencia 1", "Sentencia 2"],
        "doctrina": ["Doctrina 1", "Doctrina 2"]
    }

def cleanup_temp_file(file_path: str):
    """Remove temporary file after download"""
    try:
        os_module.unlink(file_path)
    except Exception as e:
        print(f"Error cleaning up temporary file: {e}")

@router.post("/download-document")
async def download_document(
    background_tasks: BackgroundTasks,
    document_content: dict = Body(..., description="Contenido del documento"),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """Endpoint para descargar el documento en formato Word."""
    try:
        content = document_content.get("content", "")
        document_type = document_content.get("document_type", "documento")
        
        if not content:
            raise HTTPException(status_code=400, detail="No se proporcionó contenido para el documento")
        
        # Create Word document
        doc_io = create_word_document(content, document_type)
        
        # Create a temporary file
        with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as temp_file:
            temp_file.write(doc_io.getvalue())
            temp_file_path = temp_file.name
        
        # Add cleanup task
        background_tasks.add_task(os_module.unlink, temp_file_path)
        
        # Return the file
        return FileResponse(
            temp_file_path,
            media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            filename=f"{document_type}_{datetime.now(timezone.utc).strftime('%Y%m%d')}.docx"
        )
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al generar el documento Word: {str(e)}")

@router.post(
    "/api/v1/draft-hybrid-document",
    response_model=DocumentDraftingResponse,
    summary="Generación de Documentos Híbrida",
    description="Genera documentos jurídicos aceptando tanto el formato original como el formato personalizado"
)
async def draft_hybrid_document_endpoint(
    request: HybridDocumentRequest = Body(
        ...,
        description="Parámetros para la generación del documento (formato híbrido)"
    ),
    current_user: Dict[str, Any] = Depends(get_current_user),
    http_request: Request = None
):
    """Genera documentos jurídicos con formato híbrido que acepta ambos tipos de request"""
    try:
        # Security logging
        client_ip = http_request.client.host if http_request and http_request.client else "unknown"
        security_logger.info(f"Hybrid document request from user {current_user.get('id', 'unknown')} at {client_ip}")
        
        # Determine which format is being used
        is_custom_format = (
            request.description is not None and 
            request.key_points is not None and 
            request.industry is not None and 
            request.complexity is not None
        )
        
        if is_custom_format:
            # Use custom document format
            security_logger.info("Using custom document format")
            
            # Validate and sanitize for custom format
            validated_request = validate_document_request({
                "document_type": sanitize_text_input(request.document_type),
                "description": sanitize_text_input(request.description),
                "parties": [
                    {
                        "name": sanitize_text_input(party.get("name", "")),
                        "identification": sanitize_text_input(party.get("identification", "N/A")),
                        "role": sanitize_text_input(party.get("role", ""))
                    }
                    for party in request.parties
                ],
                "key_points": sanitize_text_input(request.key_points),
                "industry": request.industry,
                "jurisdiction": request.jurisdiction,
                "complexity": request.complexity
            })
            
            # Convert parties to the format expected by the agent
            party_data = [
                {
                    "name": party["name"],
                    "identification": party["identification"],
                    "role": party["role"]
                }
                for party in validated_request["parties"]
            ]
            
            # Convert key_points string to list
            key_points_list = [point.strip() for point in validated_request["key_points"].split('\n') if point.strip()]
            
            # Security monitoring
            SecurityMonitor.log_document_generation(
                user_id=current_user.get('id', 'unknown'),
                document_type=validated_request["document_type"],
                complexity=validated_request["complexity"],
                industry=validated_request["industry"],
                client_ip=client_ip
            )
            
            result = await draft_custom_document(
                document_type=validated_request["document_type"],
                description=validated_request["description"],
                parties=party_data,
                key_points=key_points_list,
                industry=validated_request["industry"],
                jurisdiction=validated_request["jurisdiction"],
                complexity=validated_request["complexity"]
            )
            
        else:
            # Use original document format
            security_logger.info("Using original document format")
            
            # Validate and sanitize for original format
            validated_request = validate_legal_document_request({
                "document_type": request.document_type,
                "jurisdiction": request.jurisdiction,
                "key_requirements": request.key_requirements,
                "parties": request.parties,
                "legal_terms": request.legal_terms,
                "data_processing": request.data_processing
            })
            
            party_names = [party.get("name", "") for party in validated_request["parties"]]
            
            # Security monitoring
            SecurityMonitor.log_document_generation(
                user_id=current_user.get('id', 'unknown'),
                document_type=validated_request["document_type"],
                complexity="standard",  # Default for original endpoint
                industry="general",  # Default for original endpoint
                client_ip=client_ip
            )
            
            result = await draft_legal_document(
                document_type=validated_request["document_type"],
                jurisdiction=validated_request["jurisdiction"],
                key_requirements=validated_request["key_requirements"],
                parties=party_names,
                legal_terms=validated_request["legal_terms"],
                data_processing=validated_request["data_processing"]
            )
        
        # Create Word document
        doc_io = create_word_document(result["document"]["content"], result["document_type"])
        
        # Upload to R2
        document_url = await upload_document_to_r2(doc_io, result["document_type"])
        
        response_data = {
            "document": result["document"],
            "document_type": result["document_type"],
            "jurisdiction": result["jurisdiction"],
            "key_requirements": result["key_requirements"],
            "parties": result["parties"],
            "colombian_compliance": result["colombian_compliance"],
            "legal_terms": result.get("legal_terms"),
            "data_processing": result.get("data_processing"),
            "document_url": document_url,
            "timestamp": datetime.now(timezone.utc),
            "complexity": validated_request.get("complexity") if is_custom_format else "standard",
            "industry": validated_request.get("industry") if is_custom_format else "general"
        }
        
        # Success logging
        security_logger.info(f"Hybrid document generated successfully for user {current_user.get('id', 'unknown')}")
        
        return DocumentDraftingResponse(**response_data)
    
    except Exception as e:
        # Error logging
        security_logger.error(f"Error generating hybrid document for user {current_user.get('id', 'unknown')}: {str(e)}")
        SecurityMonitor.log_security_event(
            event_type="document_generation_error",
            user_id=current_user.get('id', 'unknown'),
            details={"error": str(e), "document_type": request.document_type if request else "unknown"}
        )
        
        raise HTTPException(
            status_code=500,
            detail={
                "error": str(e),
                "message": "Error en la generación del documento híbrido",
                "timestamp": datetime.now(timezone.utc).isoformat()
            }
        ) 